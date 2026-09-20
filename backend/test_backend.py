import pytest
import os
import sys

# Ensure backend directory is in sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient
from main import app
from db.database import init_db, clear_db, get_all_decisions, get_all_conflicts
from report.pdf_generator import generate_decision_passport_pdf

client = TestClient(app)

def setup_module():
    init_db()
    clear_db()

def test_root_endpoint():
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert data["name"] == "ContextLock Zero"
    assert "Git tracks code" in data["pitch"]

def test_demo_load_and_analysis():
    # 1. Trigger 1-Click Demo Load
    res = client.post("/api/demo/load")
    assert res.status_code == 200
    assert res.json()["status"] == "success"

    # 2. Verify decisions saved in SQLite
    decisions = get_all_decisions()
    assert len(decisions) >= 4
    decision_titles = [d["title"] for d in decisions]
    assert any("JWT" in t for t in decision_titles)
    assert any("OAuth" in t for t in decision_titles)
    assert any("Firebase" in t for t in decision_titles)

    # 3. Verify conflicts detected
    conflicts = get_all_conflicts()
    assert len(conflicts) >= 1
    assert any("3-Way" in c["title"] or "Conflict" in c["title"] for c in conflicts)

def test_health_endpoint():
    res = client.get("/api/health")
    assert res.status_code == 200
    health = res.json()
    assert "overall_score" in health
    assert "stability_why" in health
    assert "conflict_why" in health
    assert "ownership_why" in health
    assert health["unowned_decisions"] >= 1 # Notification architecture unassigned

def test_graph_and_impact_simulator():
    # Test Graph endpoint
    res = client.get("/api/graph")
    assert res.status_code == 200
    graph = res.json()
    assert len(graph["nodes"]) > 0
    assert len(graph["edges"]) > 0

    # Test Impact Simulator for JWT node
    jwt_node_id = graph["nodes"][0]["id"]
    sim_res = client.post("/api/simulate", json={"node_id": jwt_node_id})
    assert sim_res.status_code == 200
    sim_data = sim_res.json()
    assert sim_data["found"] is True
    assert sim_data["impacted_count"] > 0
    assert "blast_radius_score" in sim_data

def test_evidence_grounded_chat():
    res = client.post("/api/chat", json={"message": "Why did we switch authentication?"})
    assert res.status_code == 200
    chat_resp = res.json()
    assert "answer" in chat_resp
    assert "supporting_file" in chat_resp
    assert "confidence" in chat_resp
    assert len(chat_resp["citations"]) > 0

def test_decision_replay_timeline():
    res = client.get("/api/timeline")
    assert res.status_code == 200
    timeline = res.json()
    assert len(timeline["replay_milestones"]) >= 4
    assert timeline["replay_milestones"][0]["step"] == 1

def test_pdf_report_generation():
    res = client.get("/api/report")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert len(res.content) > 1000 # Valid non-empty PDF binary

def test_live_event_injection():
    res = client.post("/api/live/event", json={
        "title": "Emergency Rollback to Session Cookies",
        "description": "Due to latency overhead, the backend team initiated an emergency rollback.",
        "owner": "Sarah Chen",
        "category": "Authentication"
    })
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "success"
    assert "event_injected" in data

def test_file_listing():
    res = client.get("/api/files")
    assert res.status_code == 200
    files = res.json()
    assert isinstance(files, list)

def test_docx_file_upload_and_extraction():
    import io
    import docx
    doc = docx.Document()
    doc.add_heading("Product Requirements Document", 0)
    doc.add_paragraph("Decision: Adopt GraphQL Subscriptions for real-time telemetry updates. Owner: Sarah Chen. Date: 2025-09-16. Reason: Sub-millisecond latency for live charts.")
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()
    
    # Upload DOCX
    res = client.post(
        "/api/upload",
        files={"files": ("PRD_v2.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert res.status_code == 200
    assert res.json()["status"] == "success"

if __name__ == "__main__":
    pytest.main(["-v", "backend/test_backend.py"])
